"""
Shared Smean AI .docx style: brand tokens and the layout helpers.

Lifted verbatim from `build_literature_review.py`, which established this look
(the site's own :root palette, Georgia headings over Calibri body, teal rules
under level-1 headings, zebra tables, accent-bar callouts). That script is left
untouched and still self-contained; this module exists so a second report can
match it exactly instead of drifting.

Every builder that imports this produces a document in the same house style:

    from smean_docx import *          # tokens + helpers
    doc = new_document()
    cover(doc, title="...", subtitle="...", blurb="...", meta=[...])
    heading(doc, "1.  Something", 1, page_break=True)

Colour and typography changes belong here, not in the builders.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# --------------------------------------------------------------------------
# Smean AI brand tokens (from the site's :root CSS custom properties)
# --------------------------------------------------------------------------
TEAL = "0B776F"        # --teal        : primary
TEAL_LIGHT = "2DD4BF"  # --teal-light
TEAL_INK = "134E4A"    # --teal-ink
TEAL_SOFT = "CCFBF1"   # --teal-soft
VIOLET = "6D28D9"      # --violet      : accent
VIOLET_SOFT = "EDE9FE"
SIENNA = "C2410C"      # --sienna      : warning / caution
SIENNA_SOFT = "FFEDD5"
INK = "1A1915"         # --ink         : body text
STONE = "6B6A62"       # --stone       : muted text
OAT = "E5E1D8"         # --oat         : borders
LINEN = "F4F2EC"       # --linen       : surfaces
BONE = "FBFAF7"        # --bone        : page background
PAPER = "FFFFFF"

HEADING_FONT = "Georgia"     # brand --font-display is Fraunces, Georgia fallback
BODY_FONT = "Calibri"        # stand-in for Inter Tight (brand --font-sans)
MONO_FONT = "Courier New"   # metric-compatible everywhere; Consolas is Windows-only
KHMER_FONT = "Khmer OS System"

# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def shade(element, color):
    """Apply a solid fill to a <w:tcPr> or <w:pPr> owner."""
    pr = element.get_or_add_tcPr() if element.tag.endswith("}tc") else element.get_or_add_pPr()
    pr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": color}))


def cell_borders(cell, color=OAT, size=4, edges=("top", "bottom", "left", "right"), **overrides):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = _el("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = overrides.get(edge)
        if edge in edges or spec:
            c, s = (spec if spec else (color, size))
            borders.append(_el(f"w:{edge}", **{"w:val": "single", "w:sz": str(s),
                                               "w:space": "0", "w:color": c}))
        else:
            borders.append(_el(f"w:{edge}", **{"w:val": "nil"}))
    tcPr.append(borders)


def set_font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _el("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), KHMER_FONT if name == MONO_FONT else name)
    if size is not None:
        run.font.size = Pt(size)
        # match the complex-script size too, or Khmer runs inflate the line box
        rPr.append(_el("w:szCs", **{"w:val": str(int(round(size * 2)))}))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def inline(p, text, size=10.5, color=INK, bold_all=False, italic_all=False):
    """Render **bold**, *italic* and `code` markers into runs on paragraph `p`.

    One tokenizer shared by para-level md(), table cells and callouts, so a
    marker means the same thing everywhere. Splitting is ordered bold -> code ->
    italic; a lone `*` in prose (a glob, a footnote mark) would be mis-read, so
    write those as literal text through a chunk list instead.
    """
    for i, seg in enumerate(str(text).split("**")):
        if not seg:
            continue
        bold = bold_all or i % 2 == 1
        for j, piece in enumerate(seg.split("`")):
            if not piece:
                continue
            if j % 2 == 1:
                set_font(p.add_run(piece), MONO_FONT, size - 0.6, bold, italic_all, color)
                continue
            for k, bit in enumerate(piece.split("*")):
                if bit:
                    set_font(p.add_run(bit), BODY_FONT, size, bold,
                             italic_all or k % 2 == 1, color)


def md(doc, text, size=10.5, space_after=6, space_before=0, indent=None, align=None):
    """A body paragraph that understands **bold**, *italic* and `code`."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.28
    if indent is not None:
        pf.left_indent = Inches(indent)
    if align is not None:
        p.alignment = align
    inline(p, text, size)
    return p


def para(doc, text="", size=10.5, style=None, space_after=6, space_before=0,
         color=INK, bold=False, italic=False, align=None, font=BODY_FONT, indent=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.18
    if align is not None:
        p.alignment = align
    if indent is not None:
        pf.left_indent = Inches(indent)
    if text:
        set_font(p.add_run(text), font, size, bold, italic, color)
    return p


def rich(doc, chunks, size=10.5, space_after=6, indent=None, align=None):
    """chunks: list of (text, {bold/italic/color/font/size}) tuples."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.18
    if indent is not None:
        pf.left_indent = Inches(indent)
    if align is not None:
        p.alignment = align
    for text, opts in chunks:
        font, sz = opts.get("font", BODY_FONT), opts.get("size", size)
        bold, ital, col = opts.get("bold"), opts.get("italic"), opts.get("color", INK)
        # `code` spans inside a chunk keep the chunk's bold/italic, swap the font
        for j, piece in enumerate(str(text).split("`")):
            if piece:
                mono = j % 2 == 1
                set_font(p.add_run(piece), MONO_FONT if mono else font,
                         sz - 0.6 if mono else sz, bold, ital, col)
    return p


def heading(doc, text, level=1, page_break=False):
    if page_break:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    sizes = {1: 19, 2: 14, 3: 11.5}
    colors = {1: TEAL, 2: TEAL_INK, 3: INK}
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt({1: 4, 2: 16, 3: 12}[level])
    pf.space_after = Pt({1: 10, 2: 6, 3: 4}[level])
    pf.keep_with_next = True
    set_font(p.add_run(text), HEADING_FONT, sizes[level], True, False, colors[level])
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        bdr = _el("w:pBdr")
        bdr.append(_el("w:bottom", **{"w:val": "single", "w:sz": "12",
                                      "w:space": "6", "w:color": TEAL}))
        pPr.append(bdr)
    return p


def bullet(doc, chunks, indent=0.25, size=10.5):
    if isinstance(chunks, str):
        chunks = [(chunks, {})]
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Inches(indent + 0.2)
    pf.first_line_indent = Inches(-0.2)
    pf.space_after = Pt(3)
    pf.line_spacing = 1.15
    for text, opts in chunks:
        sz = opts.get("size", size)
        bold, ital, col = opts.get("bold"), opts.get("italic"), opts.get("color", INK)
        # inline() handles **bold**, *italic* and `code` inside the chunk; the
        # opts dict sets the baseline the markers modulate.
        inline(p, text, sz, col, bold_all=bool(bold), italic_all=bool(ital))
    return p


def table(doc, headers, rows, widths=None, size=8.8, header_fill=TEAL,
          header_color=PAPER, zebra=True, align_right=(), header=True):
    t = doc.add_table(rows=1 if header else 0, cols=len(headers))
    t.autofit = False
    t.style = "Table Grid"

    def write(cell, text, *, bold=False, color=INK, fill=None, sz=size, right=False):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.06
        if right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        inline(p, text, sz, color, bold_all=bold)
        if fill:
            shade(cell._tc, fill)
        cell_borders(cell, OAT, 4)

    if header:
        for i, h in enumerate(headers):
            write(t.rows[0].cells[i], h, bold=True, color=header_color, fill=header_fill, sz=size)
    for r, row in enumerate(rows):
        cells = t.add_row().cells
        fill = LINEN if (zebra and r % 2 == 1) else PAPER
        for i, val in enumerate(row):
            write(cells[i], val, fill=fill, right=(i in align_right))

    if widths:
        total = sum(widths)
        avail = 6.5
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(avail * w / total)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(12)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.28
    inline(p, text, 8.5, STONE, italic_all=True)
    return p


def callout(doc, title, body, accent=TEAL, fill=TEAL_SOFT):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Inches(6.5)
    shade(cell._tc, fill)
    cell_borders(cell, edges=(), left=(accent, 24), top=(fill, 4),
                 bottom=(fill, 4), right=(fill, 4))
    cell.paragraphs[0].text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.1)
    set_font(p.add_run(title), BODY_FONT, 9.5, True, False, accent)
    for line in body:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(5)
        q.paragraph_format.left_indent = Inches(0.1)
        q.paragraph_format.line_spacing = 1.15
        inline(q, line, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def code(doc, text, size=7.2):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Inches(6.5)
    shade(cell._tc, LINEN)
    cell_borders(cell, OAT, 4)
    cell.paragraphs[0].text = ""
    first = True
    for line in text.strip("\n").split("\n"):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.left_indent = Inches(0.06)
        set_font(p.add_run(line if line else " "), MONO_FONT, size, False, False, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


def field(paragraph, instr):
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    it = _el("w:instrText", **{"xml:space": "preserve"})
    it.text = instr
    r._r.append(it)
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    return r

# --------------------------------------------------------------------------
# document scaffolding
# --------------------------------------------------------------------------
ASSETS = Path(__file__).resolve().parent / "assets"
LOGO = ASSETS / "smean-logo.png"


def new_document():
    """A Document with the house Normal style and margins already applied."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.element.rPr.rFonts.set(qn("w:cs"), KHMER_FONT)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    return doc


def cover(doc, *, title, subtitle, blurb, meta, eyebrow="Speech Technology Research"):
    """The established cover: logo, eyebrow rule, title, teal subtitle, metadata."""
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    c.paragraph_format.space_before = Pt(90)
    c.paragraph_format.space_after = Pt(4)
    if LOGO.exists():
        c.add_run().add_picture(str(LOGO), width=Inches(0.82))

    rich(doc, [("SMEAN AI", {"bold": True, "size": 10, "color": TEAL}),
               ("   \u00b7   ", {"color": OAT, "size": 10}),
               (eyebrow, {"size": 10, "color": STONE})], space_after=26)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.05
    set_font(p.add_run(title), HEADING_FONT, 30, True, False, INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(subtitle), HEADING_FONT, 15, False, True, TEAL)

    hr = doc.add_paragraph()
    hr.paragraph_format.space_after = Pt(16)
    pPr = hr._p.get_or_add_pPr()
    b = _el("w:pBdr")
    b.append(_el("w:bottom", **{"w:val": "single", "w:sz": "18",
                                "w:space": "1", "w:color": TEAL}))
    pPr.append(b)

    para(doc, blurb, size=11, color=STONE, space_after=26)
    table(doc, ["", ""], meta, widths=[1.1, 3.6], size=9.5, zebra=True, header=False)


def page_footer(doc, text):
    f = doc.sections[0].footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(f.add_run(f"{text}  \u00b7  page "), BODY_FONT, 8, False, False, STONE)
    field(f, " PAGE ")
    for r in f.runs:
        set_font(r, BODY_FONT, 8, False, False, STONE)
